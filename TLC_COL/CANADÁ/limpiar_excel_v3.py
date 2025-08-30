import pandas as pd
import re
from openpyxl import load_workbook

# --- Regex patrones ---
CODE_RE = re.compile(r"^0?\d{6,10}")  # códigos arancelarios (mantiene ceros iniciales)
RATE_RE = re.compile(
    r"(\d+(\.\d+)?%|free|won/?kg|¢|/kg|/g|/ton|/l|per\s+\w+)",
    re.IGNORECASE
)

# Columnas estándar
columnas = ["Tariff Item", "Description of Goods", "Base Rate", "Staging Category"]

def parse_sheet(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        return pd.DataFrame(columns=columnas)

    max_cols = max(len(r) for r in rows)
    norm = [(r + [None]*(max_cols - len(r))) for r in rows]
    df = pd.DataFrame(norm).dropna(how="all")
    if df.empty:
        return pd.DataFrame(columns=columnas)

    # Buscar primera fila con código arancelario
    start = None
    for idx, row in df.iterrows():
        for cell in row.tolist():
            s = "" if cell is None else (str(int(cell)) if isinstance(cell, (int,float)) and float(cell).is_integer() else str(cell))
            if CODE_RE.match(s.strip()):
                start = idx
                break
        if start is not None:
            break
    if start is None:
        return pd.DataFrame(columns=columnas)

    df = df.loc[start:]

    out = []
    last_rate = None
    last_stage = None

    for _, row in df.iterrows():
        cells = []
        for c in row.tolist():
            if c is None:
                cells.append(None)
            elif isinstance(c, (int,float)) and float(c).is_integer():
                cells.append(str(int(c)))
            else:
                cells.append(str(c).strip())

        ti, desc = None, ""
        for i, c in enumerate(cells):
            if not c:
                continue
            m = CODE_RE.match(c)
            if m:
                ti = m.group(0)  # mantener el código completo con ceros
                rem = c[m.end():].strip(" -:")
                if rem:
                    desc = rem
                # Agregar celdas siguientes si no parecen tarifas
                for j in range(i+1, len(cells)):
                    cj = cells[j]
                    if not cj:
                        continue
                    if RATE_RE.search(cj) or cj in list("ABCDE"):
                        break  # detenernos, ya no es descripción
                    else:
                        desc = (desc + " " + cj).strip()
                break
        if ti is None:
            continue

        rate, stage = None, None
        for c in cells:
            if not c:
                continue
            if c in list("ABCDE"):
                stage = c
            elif RATE_RE.search(c):
                rate = c

        if rate and not stage:
            m = re.search(r"\b([ABCDE])\b\s*$", rate)
            if m:
                stage = m.group(1)
                rate = rate[:m.start()].strip()

        if rate is None:
            rate = last_rate
        else:
            last_rate = rate

        if stage is None:
            stage = last_stage
        else:
            last_stage = stage

        desc = " ".join(desc.split())
        out.append([ti, desc, rate, stage])

    res = pd.DataFrame(out, columns=columnas)
    res = res[res["Tariff Item"].astype(str).str.match(r"^\d{4,10}$")]
    return res

def process_excel(file_in, file_out):
    wb = load_workbook(file_in, data_only=True)
    all_data = []
    for name in wb.sheetnames:
        print(f"📑 Procesando hoja: {name}")
        ws = wb[name]
        df = parse_sheet(ws)
        if not df.empty:
            df.insert(0, "Source Sheet", name)
            all_data.append(df)

    if all_data:
        final = pd.concat(all_data, ignore_index=True)
        final.to_excel(file_out, index=False)
        print(f"✅ Archivo limpio guardado en: {file_out}")
    else:
        print("⚠️ No se encontró información en ninguna hoja.")

if __name__ == "__main__":
    entrada = "Lista-Productos-de-Canada.xlsx"   # archivo original
    salida = "Canada-Limpio-v2.xlsx"  # archivo limpio corregido
    process_excel(entrada, salida)
