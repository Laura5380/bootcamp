import pandas as pd
from openpyxl import load_workbook
import re

# === CONFIGURACIÓN ===
archivo_entrada = "Lista-Productos-de-Canada.xlsx"
archivo_salida = "Lista-Productos-Canada-Limpio-v2.xlsx"
columnas = ["Tariff Item", "Description of Goods", "Base Rate", "Staging Category"]

CODE_RE = re.compile(r"^(\d{4,10})\b")
RATE_RE = re.compile(r"(Free|free|%|¢|/kg|/l|/L|/ton|/t| each| per )")

def parse_sheet(ws):
    # Cargar todas las filas manteniendo None donde hay celdas combinadas
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        return pd.DataFrame(columns=columnas)

    max_cols = max(len(r) for r in rows)
    # Normalizar ancho de filas
    norm = [(r + [None]*(max_cols - len(r))) for r in rows]
    df = pd.DataFrame(norm).dropna(how="all")
    if df.empty:
        return pd.DataFrame(columns=columnas)

    # Buscar primera fila “de datos” (donde aparezca un código)
    start = None
    for idx, row in df.iterrows():
        for cell in row.tolist():
            s = "" if cell is None else (str(int(cell)) if isinstance(cell, (int,float)) and float(cell).is_integer() else str(cell))
            if CODE_RE.match(s.strip()):
                start = idx
                break
        if start is not None:
            break
    if start is None:
        return pd.DataFrame(columns=columnas)

    df = df.loc[start:]

    out = []
    last_rate = None
    last_stage = None

    for _, row in df.iterrows():
        # Normalizar a strings sencillas
        cells = []
        for c in row.tolist():
            if c is None:
                cells.append(None)
            elif isinstance(c, (int,float)) and float(c).is_integer():
                cells.append(str(int(c)))
            else:
                cells.append(str(c).strip())

        # Detectar Tariff Item + Description
        ti, desc = None, ""
        for i, c in enumerate(cells):
            if not c:
                continue
            m = CODE_RE.match(c)
            if m:
                ti = m.group(1)
                rem = c[m.end():].strip(" -:")
                if rem:
                    desc = rem
                # Sumar celdas siguientes si parecen parte de la descripción (no tarifa)
                for j in range(i+1, len(cells)):
                    cj = cells[j]
                    if not cj:
                        continue
                    if RATE_RE.search(cj) or cj in list("ABCDE"):
                        # ya no es descripción
                        pass
                    else:
                        desc = (desc + " " + cj).strip()
                break
        if ti is None:
            continue

        # Buscar Base Rate / Staging en cualquier celda de la fila
        rate, stage = None, None
        for c in cells:
            if not c:
                continue
            if c in list("ABCDE"):
                stage = c
            if RATE_RE.search(c):
                rate = c

        # Si la etapa va pegada al final de la tarifa (ej: "... 4% A")
        if rate and not stage:
            m = re.search(r"\b([ABCDE])\b\s*$", rate)
            if m:
                stage = m.group(1)
                rate = rate[:m.start()].strip()

        # Forward-fill por celdas combinadas
        if rate is None:
            rate = last_rate
        else:
            last_rate = rate

        if stage is None:
            stage = last_stage
        else:
            last_stage = stage

        # Limpieza final
        desc = " ".join(desc.split())
        out.append([ti, desc, rate, stage])

    res = pd.DataFrame(out, columns=columnas)
    res = res[res["Tariff Item"].astype(str).str.match(r"^\d{4,10}$")]
    return res

def main():
    wb = load_workbook(archivo_entrada, read_only=True)
    frames = []
    for name in wb.sheetnames:
        df = parse_sheet(wb[name])
        if not df.empty:
            df["__sheet__"] = name
            frames.append(df)

    resultado = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=columnas + ["__sheet__"])
    resultado.to_excel(archivo_salida, index=False)
    print(f"✅ Archivo limpio creado: {archivo_salida}")

if __name__ == "__main__":
    main()
